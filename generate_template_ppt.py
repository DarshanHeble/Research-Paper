import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

# Load the provided template presentation object
prs = Presentation('cu_template_ppt.pptx')

# We want to clear out any existing slides in the template to start fresh, 
# although a template usually just contains slide masters. If it has slides, we might want to keep or delete them.
# For safety, let's just append to whatever is there or assume it's an empty template.

# Slide Layouts based on inspection
TITLE_SLIDE_LAYOUT = prs.slide_layouts[0]
BULLET_SLIDE_LAYOUT = prs.slide_layouts[1]

# Title Slide
slide = prs.slides.add_slide(TITLE_SLIDE_LAYOUT)
title = slide.shapes.title
# Depending on the template, subtitle placeholder might have a different index. Let's try to find it.
subtitle = None
for shape in slide.placeholders:
    if shape != title:
        subtitle = shape
        break

if title:
    title.text = "Speech-Native Retrieval-Augmented Generation for Agricultural Advisory in Low-Resource Dialects"
if subtitle:
    subtitle.text = "Darshan Pundlik Heble\nChrist (Deemed to be University)\nBangalore, India"

# Helper to safely add text
def add_bullet_slide(title_text, bullet_points):
    slide = prs.slides.add_slide(BULLET_SLIDE_LAYOUT)
    if slide.shapes.title:
        slide.shapes.title.text = title_text
    
    # Find the body placeholder
    body = None
    for shape in slide.placeholders:
        if shape != slide.shapes.title:
            body = shape
            break
            
    if body and hasattr(body, "text_frame"):
        tf = body.text_frame
        for i, (text, level) in enumerate(bullet_points):
            if i == 0:
                tf.text = text
                tf.paragraphs[0].level = level
            else:
                p = tf.add_paragraph()
                p.text = text
                p.level = level

# Slide 2: Background & Motivation
add_bullet_slide(
    "Background & Motivation",
    [
        ("Current AI-driven agricultural advisory systems rely on a cascaded architecture:", 0),
        ("1. Automatic Speech Recognition (ASR) transcribes the spoken query.", 1),
        ("2. Retrieval-Augmented Generation (RAG) fetches information based on the text.", 1),
        ("The Flaw: ASR errors on domain-critical terms (pests, crops) propagate directly into retrieval failures.", 0)
    ]
)

# Slide 3: The Core Challenge
add_bullet_slide(
    "The Core Challenge",
    [
        ("ASR systems disproportionately fail under two compounding conditions:", 0),
        ("Domain-Critical Rare Vocabulary: Pesticide names, disease terminology, etc.", 1),
        ("Low-Resource Dialects: Colloquial and regional terms absent from standard ASR training data.", 1),
        ("Result: A mis-transcribed pest name leads to fetching irrelevant documents and generating harmful farming advice.", 0)
    ]
)

# Slide 4: Proposed Architecture
add_bullet_slide(
    "Proposed Architecture: The Four Pillars",
    [
        ("A novel framework integrating four key components:", 0),
        ("1. Speech-Native Retrieval Backbone (bypassing ASR)", 1),
        ("2. Dialect-to-Scientific-Entity Mapping Layer", 1),
        ("3. Confidence-Gated Escalation (human fallback)", 1),
        ("4. Fully Offline, Edge Deployment", 1)
    ]
)

# Slide 5: Pillar 1 & 2
add_bullet_slide(
    "Speech-Native Retrieval & Dialect Mapping",
    [
        ("Speech-Native Retrieval:", 0),
        ("Projects speech encoder representations directly into a text retriever's embedding space.", 1),
        ("Avoids transcription errors completely.", 1),
        ("Dialect-to-Scientific-Entity Mapping:", 0),
        ("Normalizes colloquial farmer terms to standard scientific nomenclature.", 1),
        ("Bootstrapped from existing public agricultural knowledge graphs (e.g., EPPO).", 1)
    ]
)

# Slide 6: Pillar 3 & 4
add_bullet_slide(
    "Confidence Gating & Offline Deployment",
    [
        ("Confidence-Gated Escalation:", 0),
        ("Incorrect advice can damage a season's crop.", 1),
        ("Evaluates internal/external confidence signals. If low, abstains and routes the query to a human expert.", 1),
        ("Offline Edge Deployment:", 0),
        ("The pipeline runs entirely offline on consumer-grade, GPU-constrained hardware (e.g., 6 GB VRAM laptop).", 1),
        ("Crucial for rural settings with limited network connectivity.", 1)
    ]
)

# Slide 7: Research Questions (RQ1-RQ4)
add_bullet_slide(
    "Research Questions",
    [
        ("RQ1: Does speech-native retrieval retain its advantage over ASR under dialect variation and rare vocabulary?", 0),
        ("RQ2: Does dialect mapping improve retrieval accuracy for colloquial queries?", 0),
        ("RQ3: Can the pipeline operate fully offline at acceptable latency/accuracy?", 0),
        ("RQ4: Does hybrid (lexical + dense) retrieval combined with dialect mapping outperform either alone on rare entities?", 0)
    ]
)

# Slide 8: Prototype Implementation & Preliminary Results
add_bullet_slide(
    "Prototype & Preliminary Results",
    [
        ("A software prototype was built and tested on synthetic, demo-scale data:", 0),
        ("Successes:", 1),
        ("Dialect mapping unlocks retrieval for colloquial text queries.", 2),
        ("Offline latency on a 6 GB VRAM laptop fits the budget (in fp16).", 2),
        ("Open Challenges:", 1),
        ("Speech-native adapter failed to generalize past its training distribution on synthetic audio.", 2),
        ("Core speech-native premise remains unconfirmed pending evaluation on real dialect speech data.", 2)
    ]
)

# Slide 9: Future Work
add_bullet_slide(
    "Conclusion & Future Work",
    [
        ("The combination of Speech-Native RAG, Dialect Mapping, Confidence Gating, and Offline Deployment is highly promising but requires further empirical validation.", 0),
        ("Future Directions:", 1),
        ("Collecting a real, recorded speech corpus for the target low-resource dialect.", 2),
        ("Improving adapter generalization in low-resource settings.", 2),
        ("Implementing dialect mapping directly in the embedding space.", 2)
    ]
)

# Save the presentation
output_path = os.path.join(os.getcwd(), 'Research_Paper_Presentation_CU.pptx')
prs.save(output_path)
print(f"Presentation saved to {output_path}")
