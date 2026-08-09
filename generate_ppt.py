import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

# Create a presentation object
prs = Presentation()

# Slide Layouts
TITLE_SLIDE_LAYOUT = prs.slide_layouts[0]
BULLET_SLIDE_LAYOUT = prs.slide_layouts[1]

# Title Slide
slide = prs.slides.add_slide(TITLE_SLIDE_LAYOUT)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Speech-Native Retrieval-Augmented Generation for Agricultural Advisory in Low-Resource Dialects"
subtitle.text = "Darshan Pundlik Heble\nChrist (Deemed to be University)\nBangalore, India"

# Slide 2: Background & Motivation
slide = prs.slides.add_slide(BULLET_SLIDE_LAYOUT)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
title.text = "Background & Motivation"
tf = body.text_frame
tf.text = "Current AI-driven agricultural advisory systems rely on a cascaded architecture:"
p = tf.add_paragraph()
p.text = "1. Automatic Speech Recognition (ASR) transcribes the spoken query."
p.level = 1
p = tf.add_paragraph()
p.text = "2. Retrieval-Augmented Generation (RAG) fetches information based on the text."
p.level = 1
p = tf.add_paragraph()
p.text = "The Flaw: ASR errors on domain-critical terms (pests, crops) propagate directly into retrieval failures."
p.level = 0

# Slide 3: The Core Challenge
slide = prs.slides.add_slide(BULLET_SLIDE_LAYOUT)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
title.text = "The Core Challenge"
tf = body.text_frame
tf.text = "ASR systems disproportionately fail under two compounding conditions:"
p = tf.add_paragraph()
p.text = "Domain-Critical Rare Vocabulary: Pesticide names, disease terminology, etc."
p.level = 1
p = tf.add_paragraph()
p.text = "Low-Resource Dialects: Colloquial and regional terms absent from standard ASR training data."
p.level = 1
p = tf.add_paragraph()
p.text = "Result: A mis-transcribed pest name leads to fetching irrelevant documents and generating harmful farming advice."
p.level = 0

# Slide 4: Proposed Architecture
slide = prs.slides.add_slide(BULLET_SLIDE_LAYOUT)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
title.text = "Proposed Architecture: The Four Pillars"
tf = body.text_frame
tf.text = "A novel framework integrating four key components:"
p = tf.add_paragraph()
p.text = "1. Speech-Native Retrieval Backbone (bypassing ASR)"
p.level = 1
p = tf.add_paragraph()
p.text = "2. Dialect-to-Scientific-Entity Mapping Layer"
p.level = 1
p = tf.add_paragraph()
p.text = "3. Confidence-Gated Escalation (human fallback)"
p.level = 1
p = tf.add_paragraph()
p.text = "4. Fully Offline, Edge Deployment"
p.level = 1

# Slide 5: Pillar 1 & 2
slide = prs.slides.add_slide(BULLET_SLIDE_LAYOUT)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
title.text = "Speech-Native Retrieval & Dialect Mapping"
tf = body.text_frame
tf.text = "Speech-Native Retrieval:"
p = tf.add_paragraph()
p.text = "Projects speech encoder representations directly into a text retriever's embedding space."
p.level = 1
p = tf.add_paragraph()
p.text = "Avoids transcription errors completely."
p.level = 1
p = tf.add_paragraph()
p.text = "Dialect-to-Scientific-Entity Mapping:"
p.level = 0
p = tf.add_paragraph()
p.text = "Normalizes colloquial farmer terms to standard scientific nomenclature."
p.level = 1
p = tf.add_paragraph()
p.text = "Bootstrapped from existing public agricultural knowledge graphs (e.g., EPPO)."
p.level = 1

# Slide 6: Pillar 3 & 4
slide = prs.slides.add_slide(BULLET_SLIDE_LAYOUT)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
title.text = "Confidence Gating & Offline Deployment"
tf = body.text_frame
tf.text = "Confidence-Gated Escalation:"
p = tf.add_paragraph()
p.text = "Incorrect advice can damage a season's crop."
p.level = 1
p = tf.add_paragraph()
p.text = "Evaluates internal/external confidence signals. If low, abstains and routes the query to a human expert."
p.level = 1
p = tf.add_paragraph()
p.text = "Offline Edge Deployment:"
p.level = 0
p = tf.add_paragraph()
p.text = "The pipeline runs entirely offline on consumer-grade, GPU-constrained hardware (e.g., 6 GB VRAM laptop)."
p.level = 1
p = tf.add_paragraph()
p.text = "Crucial for rural settings with limited network connectivity."
p.level = 1

# Slide 7: Research Questions (RQ1-RQ4)
slide = prs.slides.add_slide(BULLET_SLIDE_LAYOUT)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
title.text = "Research Questions"
tf = body.text_frame
tf.text = "RQ1: Does speech-native retrieval retain its advantage over ASR under dialect variation and rare vocabulary?"
p = tf.add_paragraph()
p.text = "RQ2: Does dialect mapping improve retrieval accuracy for colloquial queries?"
p = tf.add_paragraph()
p.text = "RQ3: Can the pipeline operate fully offline at acceptable latency/accuracy?"
p = tf.add_paragraph()
p.text = "RQ4: Does hybrid (lexical + dense) retrieval combined with dialect mapping outperform either alone on rare entities?"

# Slide 8: Prototype Implementation & Preliminary Results
slide = prs.slides.add_slide(BULLET_SLIDE_LAYOUT)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
title.text = "Prototype & Preliminary Results"
tf = body.text_frame
tf.text = "A software prototype was built and tested on synthetic, demo-scale data:"
p = tf.add_paragraph()
p.text = "Successes:"
p.level = 1
p = tf.add_paragraph()
p.text = "Dialect mapping unlocks retrieval for colloquial text queries."
p.level = 2
p = tf.add_paragraph()
p.text = "Offline latency on a 6 GB VRAM laptop fits the budget (in fp16)."
p.level = 2
p = tf.add_paragraph()
p.text = "Open Challenges:"
p.level = 1
p = tf.add_paragraph()
p.text = "Speech-native adapter failed to generalize past its training distribution on synthetic audio."
p.level = 2
p = tf.add_paragraph()
p.text = "Core speech-native premise remains unconfirmed pending evaluation on real dialect speech data."
p.level = 2

# Slide 9: Future Work
slide = prs.slides.add_slide(BULLET_SLIDE_LAYOUT)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
title.text = "Conclusion & Future Work"
tf = body.text_frame
tf.text = "The combination of Speech-Native RAG, Dialect Mapping, Confidence Gating, and Offline Deployment is highly promising but requires further empirical validation."
p = tf.add_paragraph()
p.text = "Future Directions:"
p.level = 1
p = tf.add_paragraph()
p.text = "Collecting a real, recorded speech corpus for the target low-resource dialect."
p.level = 2
p = tf.add_paragraph()
p.text = "Improving adapter generalization in low-resource settings."
p.level = 2
p = tf.add_paragraph()
p.text = "Implementing dialect mapping directly in the embedding space."
p.level = 2

# Save the presentation
output_path = os.path.join(os.getcwd(), 'Research_Paper_Presentation.pptx')
prs.save(output_path)
print(f"Presentation saved to {output_path}")
